"""Test-only helpers that build small PDF/DOCX/video fixtures on the
fly, so the repo doesn't need to carry binary sample files around."""

from __future__ import annotations

import subprocess
from pathlib import Path


def make_pdf(path: Path, paragraphs: list[str], image_count: int = 0) -> Path:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    rect = pymupdf.Rect(72, 72, page.rect.width - 72, page.rect.height - 72)
    page.insert_textbox(rect, "\n\n".join(paragraphs), fontsize=11)

    for i in range(image_count):
        png_path = path.parent / f"_fixture_img_{i}.png"
        _make_png(png_path, color=(200, 50 * i % 255, 100))
        page.insert_image(pymupdf.Rect(72 + i * 40, page.rect.height - 150, 112 + i * 40, page.rect.height - 110), filename=str(png_path))

    doc.save(str(path))
    doc.close()
    return path


def make_docx(path: Path, paragraphs: list[str], image_count: int = 0) -> Path:
    import docx

    document = docx.Document()
    for para in paragraphs:
        document.add_paragraph(para)
    for i in range(image_count):
        png_path = path.parent / f"_fixture_img_{i}.png"
        _make_png(png_path, color=(50, 200, 50 * i % 255))
        document.add_picture(str(png_path))
    document.save(str(path))
    return path


def _make_png(path: Path, color=(255, 0, 0), size=(64, 64)) -> Path:
    from PIL import Image

    img = Image.new("RGB", size, color=color)
    img.save(path)
    return path


def make_two_scene_video(path: Path, seg_seconds: float = 2.0, fps: int = 10, with_speech: bool = False) -> Path:
    """A short silent (or narrated) video with a hard color-change scene
    cut halfway through, for testing frame-difference scene detection."""
    tmp_dir = path.parent
    seg1 = tmp_dir / "_seg1.mp4"
    seg2 = tmp_dir / "_seg2.mp4"
    concat_list = tmp_dir / "_concat.txt"

    _run_ffmpeg([
        "-f", "lavfi", "-i", f"color=c=red:s=320x240:d={seg_seconds}:r={fps}",
        "-pix_fmt", "yuv420p", str(seg1),
    ])
    _run_ffmpeg([
        "-f", "lavfi", "-i", f"color=c=blue:s=320x240:d={seg_seconds}:r={fps}",
        "-pix_fmt", "yuv420p", str(seg2),
    ])
    concat_list.write_text(f"file '{seg1.name}'\nfile '{seg2.name}'\n")

    silent_video = tmp_dir / "_silent.mp4"
    _run_ffmpeg([
        "-f", "concat", "-safe", "0", "-i", str(concat_list),
        "-c", "copy", str(silent_video),
    ], cwd=tmp_dir)

    if not with_speech:
        silent_video.replace(path)
        return path

    speech_wav = tmp_dir / "_speech.wav"
    subprocess.run(
        ["espeak-ng", "-w", str(speech_wav),
         "This project implements a bubble sort algorithm that sorts a list of numbers."],
        check=True, capture_output=True,
    )
    _run_ffmpeg([
        "-i", str(silent_video), "-i", str(speech_wav),
        "-c:v", "copy", "-c:a", "aac", "-shortest", str(path),
    ])
    return path


def _run_ffmpeg(args: list[str], cwd: Path | None = None) -> None:
    subprocess.run(
        ["ffmpeg", "-y", "-loglevel", "error", *args],
        check=True, capture_output=True, cwd=cwd,
    )

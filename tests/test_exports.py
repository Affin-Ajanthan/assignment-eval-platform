"""Excel downloads: POST /exports/{code-similarity,report-similarity,batch-evaluation}.xlsx."""

import io
import sys
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.main import app as fastapi_app
from tests.auth_helpers import auth_headers, setup_subject_with_users

client = TestClient(fastapi_app)
H = auth_headers
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _code_result(subject_id, student_a="Alice"):
    return {
        "subject_id": subject_id, "assignment_name": "Lab 1: Grades", "compared": 3,
        "pairs": [
            {"submission_a_id": 1, "submission_b_id": 2, "student_a": student_a, "student_b": "Bob",
             "shared_fingerprints": 40, "jaccard": 0.82, "containment": 0.9, "semantic_similarity": 0.95,
             "token_flagged": True, "flagged": True, "flag_level": "high", "flag_reason": "High token overlap"},
            {"submission_a_id": 1, "submission_b_id": 3, "student_a": student_a, "student_b": "Chamath",
             "shared_fingerprints": 3, "jaccard": 0.05, "containment": 0.1, "semantic_similarity": None,
             "token_flagged": False, "flagged": False, "flag_level": None, "flag_reason": None},
        ],
        "semantic": {"status": "ok", "model": "microsoft/unixcoder-base-nine", "warning": None,
                     "submissions_encoded": 3, "cache_hits": 0, "skipped_too_small": [], "pairs_compared": 2,
                     "cohort_median": 0.7, "review_threshold": 0.8},
        "starter_code_excluded": True,
    }


def _report_result(subject_id):
    return {
        "subject_id": subject_id, "assignment_name": "Lab 1: Grades", "compared": 2,
        "skipped": [{"submission_id": 9, "student": "Dilini", "reason": "No report submitted"}],
        "common_fingerprints_excluded": 4,
        "pairs": [{"submission_a_id": 1, "submission_b_id": 2, "student_a": "Alice", "student_b": "Bob",
                   "words_a": 900, "words_b": 950, "shared_fingerprints": 30, "jaccard": 0.35, "containment": 0.5,
                   "coverage_a": 0.6, "coverage_b": 0.55, "semantic_similarity": 0.9, "flag_level": "high",
                   "flag_reason": "Identical passages",
                   "passages": [{"words": 120, "text_a": "copied text", "text_b": "copied text"}]}],
        "semantic": {"status": "ok", "model": "sentence-transformers/all-MiniLM-L6-v2", "warning": None,
                     "pairs_compared": 1, "cohort_median": 0.7, "review_threshold": 0.85},
        "exploratory": True,
    }


def _batch_result(subject_id):
    return {
        "subject_id": subject_id, "assignment_name": "Lab 1: Grades", "rubric_id": 999999,
        "total": 3, "evaluated": 2, "skipped": 0, "failed": 1,
        "results": [
            {"submission_id": 1, "student_name": "Alice", "status": "evaluated", "recommended_score": 8,
             "recommended_max": 10, "grader": "heuristic", "review_flags": [], "error": None},
            {"submission_id": 2, "student_name": "Bob", "status": "evaluated", "recommended_score": 5,
             "recommended_max": 10, "grader": "heuristic", "review_flags": ["Similarity flag"], "error": None},
            {"submission_id": 3, "student_name": "Chamath", "status": "failed", "recommended_score": None,
             "recommended_max": None, "grader": None, "review_flags": [], "error": "RuntimeError: boom"},
        ],
    }


def _workbook(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"].startswith("attachment;")
    assert ".xlsx" in response.headers["content-disposition"]
    return load_workbook(io.BytesIO(response.content))


def test_code_similarity_export():
    subject, tok, _ = setup_subject_with_users(client)
    wb = _workbook(client.post("/exports/code-similarity.xlsx", json=_code_result(subject["id"]), headers=H(tok)))
    assert wb.sheetnames == ["Summary", "Pairs"]
    pairs = wb["Pairs"]
    assert [c.value for c in pairs[1]][:3] == ["Student A", "Student B", "Submission A"]
    assert pairs["A2"].value == "Alice" and pairs["B2"].value == "Bob"  # flagged pair first
    assert pairs["F2"].value == 0.82 and pairs["F2"].number_format == "0%"
    assert pairs["H3"].value is None  # no semantic score
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert summary["Flagged pairs"] == 1 and summary["Starter code excluded"] == "Yes"
    assert subject["name"] in summary["Subject"]


def test_report_similarity_export_includes_passages_and_skipped():
    subject, tok, _ = setup_subject_with_users(client)
    wb = _workbook(client.post("/exports/report-similarity.xlsx", json=_report_result(subject["id"]),
                               headers=H(tok)))
    assert wb.sheetnames == ["Summary", "Pairs", "Passages", "Not compared"]
    assert wb["Pairs"]["E2"].value == 0.6
    assert wb["Passages"]["D2"].value == 120 and wb["Passages"]["E2"].value == "copied text"
    assert wb["Not compared"]["C2"].value == "No report submitted"


def test_batch_evaluation_export():
    subject, tok, _ = setup_subject_with_users(client)
    wb = _workbook(client.post("/exports/batch-evaluation.xlsx", json=_batch_result(subject["id"]), headers=H(tok)))
    results = wb["Results"]
    assert [results.cell(row=r, column=2).value for r in (2, 3, 4)] == ["Alice", "Bob", "Chamath"]
    assert results["F2"].value == 0.8 and results["F2"].number_format == "0%"
    assert results["H3"].value == "Similarity flag"
    assert results["C4"].value == "Failed" and "boom" in results["I4"].value
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert summary["Rubric"] == "Rubric #999999" and summary["Average recommended score"] == "65%"


def test_formula_like_names_are_written_as_text():
    subject, tok, _ = setup_subject_with_users(client)
    evil = '=HYPERLINK("http://example.com","click")'
    wb = _workbook(client.post("/exports/code-similarity.xlsx", json=_code_result(subject["id"], student_a=evil),
                               headers=H(tok)))
    cell = wb["Pairs"]["A2"]
    assert cell.data_type == "s" and cell.value == evil


def test_exports_are_instructor_only_and_subject_scoped():
    subject, tok, student_tok = setup_subject_with_users(client)
    _, other_tok, _ = setup_subject_with_users(client)
    for path, body in [("/exports/code-similarity.xlsx", _code_result(subject["id"])),
                       ("/exports/report-similarity.xlsx", _report_result(subject["id"])),
                       ("/exports/batch-evaluation.xlsx", _batch_result(subject["id"]))]:
        assert client.post(path, json=body, headers=H(student_tok)).status_code == 403
        assert client.post(path, json=body, headers=H(other_tok)).status_code == 403
        assert client.post(path, json=body).status_code == 401
    assert client.post("/exports/code-similarity.xlsx", json={"subject_id": subject["id"]},
                       headers=H(tok)).status_code == 422

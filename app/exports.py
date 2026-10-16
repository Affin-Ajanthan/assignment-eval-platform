"""Excel (.xlsx) exports of the instructor's evaluation results.

Built from the results the instructor already has on screen (the similarity
check, the report similarity check and "Evaluate a whole assignment"), so
nothing is re-run. Uses openpyxl.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
_HIGH_FILL = PatternFill("solid", fgColor="F8D7DA")    # identical code/text
_REVIEW_FILL = PatternFill("solid", fgColor="FFF3CD")  # semantic only
_FAILED_FILL = PatternFill("solid", fgColor="F8D7DA")
_PERCENT = "0%"
_NOTE = ("Flags are review signals for the instructor, not proof of misconduct. "
         "Thresholds have not been calibrated on real submissions.")


def _write_row(ws, values, *, bold=False, fill=None, formats=None):
    ws.append(list(values))
    row = ws.max_row
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col)
        if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
            # Text that looks like a formula stays text, so a student name
            # such as '=HYPERLINK(...)' can't run when the file is opened.
            cell.data_type = "s"
        if bold:
            cell.font = Font(bold=True)
        if fill is not None:
            cell.fill = fill
        if formats and formats.get(col):
            cell.number_format = formats[col]
        cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 60)


def _table(ws, header, rows, *, formats=None, fills=None, widths=None):
    _write_row(ws, header, bold=True, fill=_HEADER_FILL)
    ws.freeze_panes = "A2"
    for i, row in enumerate(rows):
        _write_row(ws, row, fill=fills[i] if fills else None, formats=formats)
    for col, title in enumerate(header, start=1):
        if widths and col <= len(widths) and widths[col - 1]:
            width = widths[col - 1]
        else:
            longest = max([len(str(title))] + [len(str(r[col - 1])) for r in rows if r[col - 1] is not None])
            width = min(max(10, longest + 2), 60)
        ws.column_dimensions[get_column_letter(col)].width = width
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"


def _summary(ws, title, items):
    ws.title = "Summary"
    _write_row(ws, [title], bold=True)
    ws.append([])
    for label, value in items:
        _write_row(ws, [label, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])
    _write_row(ws, [_NOTE])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60


def _bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _generated() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _flag_fill(level):
    return _HIGH_FILL if level == "high" else _REVIEW_FILL if level == "review" else None


def _flag_label(level):
    return {"high": "Review (identical content)", "review": "Review (similar meaning)"}.get(level, "")


def filename(kind: str, assignment_name: str) -> str:
    """e.g. code-similarity_Assignment-3_2026-10-13.xlsx (ASCII only)."""
    slug = re.sub(r"[^A-Za-z0-9]+", "-", assignment_name).strip("-")[:60] or "assignment"
    return f"{kind}_{slug}_{datetime.now(timezone.utc):%Y-%m-%d}.xlsx"


# --------------------------------------------------------------------------
# The three exports
# --------------------------------------------------------------------------

def code_similarity_xlsx(result, subject_name: str) -> bytes:
    """Similarity check (Winnowing + UniXcoder) -> Summary + Pairs."""
    wb = Workbook()
    sem = result.semantic
    _summary(wb.active, "Code similarity check", [
        ("Subject", subject_name),
        ("Assignment", result.assignment_name),
        ("Submissions compared", result.compared),
        ("Pairs", len(result.pairs)),
        ("Flagged pairs", sum(1 for p in result.pairs if p.flagged)),
        ("Starter code excluded", "Yes" if result.starter_code_excluded else "No"),
        ("Semantic model", f"{sem.model} ({sem.status})"),
        ("Semantic cohort median", sem.cohort_median if sem.cohort_median is not None else "-"),
        ("Semantic review threshold", sem.review_threshold if sem.review_threshold is not None else "-"),
        ("Semantic warning", sem.warning or "-"),
        ("Exported", _generated()),
    ])
    pairs = sorted(result.pairs, key=lambda p: (not p.flagged, -p.containment))
    rows = [[p.student_a, p.student_b, p.submission_a_id, p.submission_b_id, p.shared_fingerprints,
             p.jaccard, p.containment, p.semantic_similarity, _flag_label(p.flag_level), p.flag_reason or ""]
            for p in pairs]
    _table(wb.create_sheet("Pairs"),
           ["Student A", "Student B", "Submission A", "Submission B", "Shared fingerprints", "Jaccard",
            "Containment", "Semantic (UniXcoder)", "Flag", "Reason"],
           rows, formats={6: _PERCENT, 7: _PERCENT, 8: _PERCENT},
           fills=[_flag_fill(p.flag_level) for p in pairs], widths=[22, 22, 13, 13, 12, 10, 12, 12, 26, 50])
    return _bytes(wb)


def report_similarity_xlsx(result, subject_name: str) -> bytes:
    """Report similarity (PDF) -> Summary + Pairs + Passages + Not compared."""
    wb = Workbook()
    sem = result.semantic
    _summary(wb.active, "Report similarity check (PDF)", [
        ("Subject", subject_name),
        ("Assignment", result.assignment_name),
        ("Reports compared", result.compared),
        ("Reports not compared", len(result.skipped)),
        ("Pairs", len(result.pairs)),
        ("Flagged pairs", sum(1 for p in result.pairs if p.flag_level)),
        ("Fingerprints shared by most reports (ignored)", result.common_fingerprints_excluded),
        ("Template PDF cut out of reports", result.template_filename or "None"),
        ("Template words cut (all reports)", sum(result.template_words_removed.values())),
        ("Semantic model", f"{sem.model} ({sem.status})"),
        ("Semantic cohort median", sem.cohort_median if sem.cohort_median is not None else "-"),
        ("Semantic review threshold", sem.review_threshold if sem.review_threshold is not None else "-"),
        ("Semantic warning", sem.warning or "-"),
        ("Exported", _generated()),
    ])
    pairs = sorted(result.pairs, key=lambda p: (p.flag_level is None, -max(p.coverage_a, p.coverage_b)))
    rows = [[p.student_a, p.student_b, p.words_a, p.words_b, p.coverage_a, p.coverage_b, p.jaccard,
             p.containment, p.semantic_similarity, len(p.passages), _flag_label(p.flag_level), p.flag_reason or ""]
            for p in pairs]
    _table(wb.create_sheet("Pairs"),
           ["Student A", "Student B", "Words A", "Words B", "Identical text A", "Identical text B", "Jaccard",
            "Containment", "Semantic (MiniLM)", "Passages", "Flag", "Reason"],
           rows, formats={5: _PERCENT, 6: _PERCENT, 7: _PERCENT, 8: _PERCENT, 9: _PERCENT},
           fills=[_flag_fill(p.flag_level) for p in pairs], widths=[22, 22, 9, 9, 12, 12, 9, 12, 12, 9, 26, 50])
    passage_rows = [[p.student_a, p.student_b, i, x.words, x.text_a, x.text_b]
                    for p in pairs for i, x in enumerate(p.passages, 1)]
    _table(wb.create_sheet("Passages"), ["Student A", "Student B", "#", "Identical words", "Text in A", "Text in B"],
           passage_rows, widths=[22, 22, 5, 10, 70, 70])
    _table(wb.create_sheet("Not compared"), ["Submission", "Student", "Reason"],
           [[s.submission_id, s.student, s.reason] for s in result.skipped], widths=[12, 26, 70])
    return _bytes(wb)


def _flag_text(flagged):
    """Flag column text; blank when unknown (a failed submission)."""
    return "" if flagged is None else "Flagged" if flagged else "No flag"


def batch_evaluation_xlsx(result, subject_name: str, rubric_name: str) -> bytes:
    """Evaluate a whole assignment -> Summary + Results."""
    wb = Workbook()
    scores = [r.recommended_score / r.recommended_max for r in result.results
              if r.status == "evaluated" and r.recommended_score is not None and r.recommended_max]
    _summary(wb.active, "Automated evaluation of a whole assignment", [
        ("Subject", subject_name),
        ("Assignment", result.assignment_name),
        ("Rubric", rubric_name),
        ("Submissions", result.total),
        ("Evaluated", result.evaluated),
        ("Skipped (already evaluated)", result.skipped),
        ("Failed", result.failed),
        ("Average recommended score", f"{sum(scores) / len(scores):.0%}" if scores else "-"),
        ("Exported", _generated()),
    ])
    rows, fills = [], []
    for r in result.results:
        pct = (r.recommended_score / r.recommended_max
               if r.recommended_score is not None and r.recommended_max else None)
        rows.append([r.submission_id, r.student_name, r.status.capitalize(), r.recommended_score,
                     r.recommended_max, pct, _flag_text(r.similarity_flagged), _flag_text(r.ai_flagged),
                     r.grader or "", "; ".join(r.review_flags), r.error or ""])
        fills.append(_FAILED_FILL if r.status == "failed" else _REVIEW_FILL if r.review_flags else None)
    _table(wb.create_sheet("Results"),
           ["Submission", "Student", "Status", "Recommended score", "Out of", "Percent",
            "Similarity flag", "AI detection flag", "Grader", "Review flags", "Error"],
           rows, formats={6: _PERCENT}, fills=fills,
           widths=[12, 24, 11, 12, 9, 9, 15, 17, 12, 60, 40])
    ws = wb["Results"]
    ws.append([])
    ws.append(["Automated scores are a second opinion; the instructor's manual grade always counts."])
    return _bytes(wb)

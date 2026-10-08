"""
Report pipeline: text/image extraction from a submitted report, an
AI-generated-text style heuristic, and a screenshot/claim sanity check.

Mirrors three nodes from the platform plan's pipeline diagram:
  E1 text extraction, E2 AI-text detection, E3 image/screenshot check.

As with the code AI-heuristic in the similarity checker, E2 here is a
style heuristic, not a validated classifier (GPTZero/Originality.ai are
paid APIs; this module gives you a free, offline first pass and a place
to plug a real API in later -- see ``score_text`` docstring).
"""

from __future__ import annotations

import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path

# --------------------------------------------------------------------------
# Extraction (E1)
# --------------------------------------------------------------------------


@dataclass
class ReportContent:
    text: str
    image_count: int
    unit_count: int  # pages (PDF) or paragraphs (DOCX/text)
    source_format: str


def _extract_pdf(path: Path) -> ReportContent:
    import pymupdf

    doc = pymupdf.open(path)
    text_parts = []
    image_count = 0
    for page in doc:
        text_parts.append(page.get_text())
        image_count += len(page.get_images(full=True))
    page_count = doc.page_count
    doc.close()
    return ReportContent(
        text="\n".join(text_parts),
        image_count=image_count,
        unit_count=page_count,
        source_format="pdf",
    )


def _extract_docx(path: Path) -> ReportContent:
    import docx

    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    # Inline images live in the document's related parts as image blobs.
    image_count = sum(
        1 for rel in document.part.rels.values() if "image" in rel.reltype
    )
    return ReportContent(
        text=text,
        image_count=image_count,
        unit_count=len(document.paragraphs),
        source_format="docx",
    )


def _extract_plain_text(path: Path) -> ReportContent:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return ReportContent(
        text=text,
        image_count=0,
        unit_count=len([p for p in text.split("\n\n") if p.strip()]),
        source_format=path.suffix.lstrip(".") or "text",
    )


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def extract_report(path: Path) -> ReportContent:
    """Extract text and count embedded images from a report file.
    Supports PDF and DOCX; falls back to plain-text reading for
    anything else (.md, .txt) so the pipeline never hard-fails on an
    unexpected format."""
    path = Path(path)
    extractor = _EXTRACTORS.get(path.suffix.lower(), _extract_plain_text)
    return extractor(path)


# --------------------------------------------------------------------------
# AI-text style heuristic (E2)
# --------------------------------------------------------------------------

_CLICHE_PHRASES = [
    "in conclusion", "it is important to note", "furthermore",
    "overall, it can be seen", "in today's world", "delve into",
    "it is worth noting", "in summary", "moreover", "as an ai language model",
    "plays a crucial role", "in the realm of", "on the other hand",
    "it is essential to", "this comprehensive", "let's dive in",
]

_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?])\s+")
_WORD_RE = re.compile(r"[A-Za-z']+")


@dataclass
class TextHeuristicScore:
    score: int  # 0-100, higher = more AI-like signal
    signal: str  # "low" | "medium" | "high"
    reasons: list[str] = field(default_factory=list)
    # "style-heuristic" (score_text below) or "fast-detectgpt:<model>"
    # (ai_text_detection.py, used when its model is available).
    method: str = "style-heuristic"


def score_text(text: str) -> TextHeuristicScore:
    """Experimental AI-generated-text style score.

    Real deployments should call a maintained classifier (GPTZero,
    Originality.ai, etc. -- see the platform plan's tech-stack table)
    behind this same signature: swap the body of this function for an
    API call and every caller (rubric_grading, pipeline) keeps working
    unchanged. This offline version needs no API key or network access,
    at the cost of being far less accurate.
    """
    sentences = [s for s in _SENTENCE_SPLIT_RE.split(text) if s.strip()]
    words = _WORD_RE.findall(text.lower())

    reasons: list[str] = []
    points = 0

    if len(sentences) >= 5:
        lengths = [len(_WORD_RE.findall(s)) for s in sentences]
        mean_len = statistics.mean(lengths)
        stdev_len = statistics.pstdev(lengths)
        # Low sentence-length variance ("burstiness") is a commonly cited
        # signal of machine-generated prose (humans vary sentence length more).
        if mean_len > 0 and (stdev_len / mean_len) < 0.35:
            points += 30
            reasons.append("sentence lengths are unusually uniform (low burstiness)")

    if len(words) >= 50:
        ttr = len(set(words)) / len(words)
        if ttr < 0.4:
            points += 25
            reasons.append(f"low vocabulary diversity (type-token ratio {ttr:.2f})")

    lower_text = text.lower()
    hits = [p for p in _CLICHE_PHRASES if p in lower_text]
    if hits:
        points += min(30, 10 * len(hits))
        reasons.append(f"contains common LLM stock phrases: {', '.join(hits[:3])}")

    if words:
        avg_word_len = sum(len(w) for w in words) / len(words)
        if avg_word_len > 5.2:
            points += 15
            reasons.append("unusually long average word length")

    points = min(points, 100)
    signal = "low" if points < 34 else "medium" if points < 67 else "high"
    if not reasons:
        reasons.append("no strong stylistic signal either way")

    return TextHeuristicScore(score=points, signal=signal, reasons=reasons)


# --------------------------------------------------------------------------
# Screenshot / claim sanity check (E3)
# --------------------------------------------------------------------------

_FIGURE_MENTION_RE = re.compile(
    r"\b(figure|fig\.?|screenshot|see below|shown below)\b", re.IGNORECASE
)


@dataclass
class ImageCheckResult:
    image_count: int
    figure_mentions: int
    flagged: bool
    reason: str


def check_images(content: ReportContent) -> ImageCheckResult:
    """Flags a report that talks about screenshots/figures it doesn't
    actually contain, or that contains images the text never refers to
    -- either can indicate a report copy-pasted from elsewhere, or one
    whose screenshots weren't actually included."""
    mentions = len(_FIGURE_MENTION_RE.findall(content.text))

    if mentions >= 2 and content.image_count == 0:
        return ImageCheckResult(
            image_count=content.image_count,
            figure_mentions=mentions,
            flagged=True,
            reason=f"text references figures/screenshots {mentions} time(s) but the report has no embedded images",
        )
    if content.image_count >= 3 and mentions == 0:
        return ImageCheckResult(
            image_count=content.image_count,
            figure_mentions=mentions,
            flagged=True,
            reason=f"report embeds {content.image_count} images but the text never refers to any of them",
        )
    return ImageCheckResult(
        image_count=content.image_count,
        figure_mentions=mentions,
        flagged=False,
        reason="image/text references are consistent",
    )


@dataclass
class ReportAnalysis:
    content: ReportContent
    ai_text: TextHeuristicScore
    image_check: ImageCheckResult


def ai_text_signal(text: str) -> TextHeuristicScore:
    """Fast-DetectGPT (ai_text_detection.py) when its local model is
    available; the style heuristic above otherwise."""
    from .ai_text_detection import get_ai_text_detector, to_score

    result = get_ai_text_detector().detect(text)
    return to_score(result) if result is not None else score_text(text)


def analyze_report(path: Path) -> ReportAnalysis:
    content = extract_report(path)
    return ReportAnalysis(
        content=content,
        ai_text=ai_text_signal(content.text),
        image_check=check_images(content),
    )
